import os
import re
import json
import shutil
import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

os.environ["HF_TOKEN"] = os.environ.get("HF_TOKEN", "")
os.environ["HF_HUB_DISABLE_SYMLINKS_WARNING"] = "1"
os.environ["TOKENIZERS_PARALLELISM"] = "false"

from fastapi import FastAPI, UploadFile, File, Form, BackgroundTasks
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse
from fastapi.staticfiles import StaticFiles
import pandas as pd
import openpyxl

import video_engine

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
IMAGES_DIR = UPLOAD_DIR / "images"
STATIC_DIR = BASE_DIR / "static"
FONTS_DIR = BASE_DIR / "fonts"
OUTPUT_DIR = BASE_DIR / "output_renders"

UPLOAD_DIR.mkdir(exist_ok=True)
IMAGES_DIR.mkdir(exist_ok=True)
STATIC_DIR.mkdir(exist_ok=True)
FONTS_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

DEFAULT_MEDIA_DIR = IMAGES_DIR if IMAGES_DIR.exists() else (BASE_DIR / "images")

app = FastAPI(title="Dhamma Studio")
app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
app.mount("/fonts", StaticFiles(directory=FONTS_DIR), name="fonts")

SESSION = {
    "audio_path": "",
    "media_dir": str(DEFAULT_MEDIA_DIR),
    "logo_path": "",
    "excel_path": "",
    "rows": [],
    "available_languages": [],
    "progress_msg": "Ready",
    "progress_pct": 0,
    "is_rendering": False,
    "last_video": "",
    "cached_images": [],
    "image_embeddings": None,
    "current_mode": "excel"
}


def set_progress(msg: str, pct: int):
    SESSION["progress_msg"] = msg
    SESSION["progress_pct"] = pct


def get_available_images():
    search_dirs = [Path(SESSION["media_dir"]), UPLOAD_DIR / "images", BASE_DIR / "images", BASE_DIR]
    all_imgs = []
    seen = set()
    for d in search_dirs:
        if d.exists() and d.is_dir():
            for p in list(d.glob("*.jpg")) + list(d.glob("*.jpeg")) + list(d.glob("*.png")) + list(d.glob("*.webp")):
                if p.name not in seen:
                    all_imgs.append(p)
                    seen.add(p.name)
    return all_imgs


def parse_time_to_seconds(val: Any) -> Optional[float]:
    if pd.isna(val) or val is None:
        return None
    if isinstance(val, (int, float)):
        return round(float(val), 2)
    if isinstance(val, (datetime.time, datetime.datetime)):
        return round(val.hour * 3600 + val.minute * 60 + val.second + val.microsecond / 1e6, 2)
    
    clean_str = str(val).strip().replace(",", ".")
    if not clean_str or clean_str == "-":
        return None

    if " - " in clean_str:
        clean_str = clean_str.split(" - ")[0].strip()

    parts = clean_str.split(":")
    try:
        if len(parts) == 1:
            return round(float(parts[0]), 2)
        elif len(parts) == 2:
            return round(float(parts[0]) * 60 + float(parts[1]), 2)
        elif len(parts) == 3:
            return round(float(parts[0]) * 3600 + float(parts[1]) * 60 + float(parts[2]), 2)
    except (ValueError, TypeError):
        return None
    return None


def parse_excel_sync_sheet(excel_path: Path, format_type: str = "landscape"):
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    sheet_names = wb.sheetnames

    target_sheet = None
    if format_type == "mobile":
        for sname in sheet_names:
            lower_name = sname.lower()
            if any(k in lower_name for k in ["mobile", "9x16", "9:16", "shorts", "2 lines"]):
                target_sheet = sname
                break
    else:
        for sname in sheet_names:
            lower_name = sname.lower()
            if any(k in lower_name for k in ["youtube", "yt", "16x9", "landscape", "4 lines"]):
                target_sheet = sname
                break

    if not target_sheet or wb[target_sheet].max_row < 2:
        target_sheet = sheet_names[0]

    raw_df = pd.read_excel(excel_path, sheet_name=target_sheet, header=None)

    header_idx = -1
    for idx, row in raw_df.iterrows():
        cells = [str(v).strip().lower() for v in row.values if pd.notna(v)]
        if len(cells) < 3:
            continue
        has_start = any(re.search(r"^(start|start_sec|စတင်ချိန်)", c) or c in ["start", "start_sec", "စတင်ချိန်"] for c in cells)
        has_end = any(re.search(r"^(end|end_sec|ပြီးဆုံးချိန်)", c) or c in ["end", "end_sec", "ပြီးဆုံးချိန်"] for c in cells)
        if has_start and has_end:
            header_idx = idx
            break

    if header_idx != -1:
        df = pd.read_excel(excel_path, sheet_name=target_sheet, header=header_idx)
    else:
        df = pd.read_excel(excel_path, sheet_name=target_sheet)

    col_start, col_end, col_index, col_manual_image = None, None, None, None

    for c in df.columns:
        c_clean = str(c).strip().lower()
        if c_clean in ["စဉ်", "no", "no.", "#", "id", "num"] or re.match(r"^(no|id|num)\b", c_clean):
            if col_index is None: col_index = c
        elif any(k in c_clean for k in ["စတင်ချိန်", "start_sec"]) or re.search(r"\b(start|start_sec)\b", c_clean):
            if col_start is None: col_start = c
        elif any(k in c_clean for k in ["ပြီးဆုံးချိန်", "end_sec"]) or re.search(r"\b(end|end_sec)\b", c_clean):
            if col_end is None: col_end = c
        elif c_clean in ["image", "img", "picture", "photo", "ပုံ"]:
            col_manual_image = c

    meta_cols = {col_start, col_end, col_index, col_manual_image}
    lang_columns = []

    for idx_col, c in enumerate(df.columns):
        if c in meta_cols:
            continue
        c_clean = str(c).strip().lower()
        if "timecode" in c_clean or "duration" in c_clean or "unnamed" in c_clean:
            continue

        m = re.match(r"^([a-zA-Z\-]+)", str(c).strip())
        lang_code = m.group(1).lower() if m else f"lang_{idx_col}"
        if "pali" in c_clean or "ပါဠိ" in c_clean:
            lang_code = "pali"

        lang_columns.append({"col_name": c, "code": lang_code, "label": str(c).strip()})

    valid_imgs = get_available_images()
    default_img_url = f"/api/get_image/{valid_imgs[0].name}" if valid_imgs else ""

    rows = []
    for _, r in df.iterrows():
        if col_start is None or col_end is None:
            continue
        st = parse_time_to_seconds(r[col_start])
        et = parse_time_to_seconds(r[col_end])
        if st is None or et is None:
            continue

        translations = {}
        for l in lang_columns:
            val = r[l["col_name"]]
            translations[l["code"]] = str(val).strip() if pd.notna(val) else ""

        pali_text = translations.get("pali", "")
        original_my = translations.get("my", "")
        main_caption = original_my or pali_text or next((v for v in translations.values() if v), "")
        if not main_caption:
            continue

        matched_img_val = default_img_url
        if col_manual_image and pd.notna(r[col_manual_image]):
            custom_img = str(r[col_manual_image]).strip()
            for cand_dir in [Path(SESSION["media_dir"]), UPLOAD_DIR / "images", BASE_DIR / "images"]:
                if (cand_dir / custom_img).exists():
                    matched_img_val = f"/api/get_image/{custom_img}"
                    break

        rows.append({
            "index": len(rows),
            "caption": main_caption,
            "original_text": original_my or main_caption,
            "pali_text": pali_text,
            "translations": translations,
            "start_time": st,
            "end_time": et,
            "duration": round(max(0.2, et - st), 2),
            "target_lang": "my" if "my" in translations else (lang_columns[0]["code"] if lang_columns else "default"),
            "matched_img": matched_img_val,
            "is_manual_img": bool(col_manual_image and pd.notna(r[col_manual_image]))
        })

    languages_meta = [{"code": l["code"], "label": l["label"]} for l in lang_columns]
    return rows, languages_meta


def auto_match_all_images():
    valid_imgs = get_available_images()
    if not valid_imgs or not SESSION["rows"]:
        return

    SESSION["cached_images"], SESSION["image_embeddings"] = video_engine.scan_and_index_images(
        valid_imgs, progress_cb=set_progress
    )

    if not SESSION["cached_images"] or SESSION["image_embeddings"] is None:
        return

    clip_query_texts, indices = [], []
    for i, r in enumerate(SESSION["rows"]):
        if not r.get("is_manual_img"):
            trans = r.get("translations", {})
            en_text = trans.get("en") or trans.get("english") or r.get("original_text", "")
            clean_en = re.sub(r"[\r\n]+", " ", str(en_text)).strip()
            if clean_en and clean_en.lower() != "nan":
                clip_query_texts.append(clean_en)
                indices.append(i)

    if clip_query_texts:
        matched = video_engine.batch_match_images_hybrid(
            clip_query_texts,
            SESSION["cached_images"],
            SESSION["image_embeddings"],
            top_k=len(SESSION["cached_images"]),
            penalty_weight=0.35
        )
        for row_idx, img_p in zip(indices, matched):
            if img_p and img_p.name:
                SESSION["rows"][row_idx]["matched_img"] = f"/api/get_image/{img_p.name}"


@app.get("/", response_class=HTMLResponse)
async def serve_index():
    with open(STATIC_DIR / "index.html", "r", encoding="utf-8") as f:
        return f.read()


@app.get("/api/status")
async def get_status():
    return {
        "msg": SESSION["progress_msg"],
        "pct": SESSION["progress_pct"],
        "is_rendering": SESSION["is_rendering"],
        "total_rows": len(SESSION["rows"]),
        "has_video": bool(SESSION["last_video"] and os.path.exists(SESSION["last_video"]))
    }


@app.post("/api/upload_images_folder")
async def upload_images_folder(files: List[UploadFile] = File(...)):
    IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    saved = []
    for f in files:
        if f.filename:
            dest = IMAGES_DIR / Path(f.filename).name
            with open(dest, "wb") as buffer:
                shutil.copyfileobj(f.file, buffer)
            if dest.suffix.lower() in [".jpg", ".jpeg", ".png", ".webp"]:
                saved.append(dest)

    SESSION["media_dir"] = str(IMAGES_DIR)
    if SESSION["rows"]:
        auto_match_all_images()
    set_progress(f"{len(saved)} Images Ready", 100)
    return {"success": True, "count": len(saved), "rows": SESSION["rows"]}


@app.post("/api/match_images")
async def match_images():
    if not SESSION["rows"]:
        return JSONResponse({"success": False, "error": "Timeline is empty."}, status_code=400)
    auto_match_all_images()
    set_progress("Images matched!", 100)
    return {"success": True, "rows": SESSION["rows"]}


@app.post("/api/upload_excel_sync")
async def upload_excel_sync(
    audio: UploadFile = File(...),
    excel: UploadFile = File(...),
    logo: Optional[UploadFile] = File(None),
    format: str = Form("landscape")
):
    try:
        set_progress("Processing Excel...", 20)
        audio_path = UPLOAD_DIR / audio.filename
        with open(audio_path, "wb") as f:
            shutil.copyfileobj(audio.file, f)
        SESSION["audio_path"] = str(audio_path)

        excel_path = UPLOAD_DIR / excel.filename
        with open(excel_path, "wb") as f:
            shutil.copyfileobj(excel.file, f)
        SESSION["excel_path"] = str(excel_path)

        if logo and logo.filename:
            logo_path = UPLOAD_DIR / logo.filename
            with open(logo_path, "wb") as f:
                shutil.copyfileobj(logo.file, f)
            SESSION["logo_path"] = str(logo_path)
        else:
            SESSION["logo_path"] = ""

        rows, languages = parse_excel_sync_sheet(excel_path, format_type=format)
        if not rows:
            return JSONResponse({"success": False, "error": "No valid timestamps found."}, status_code=400)

        SESSION["rows"] = rows
        SESSION["available_languages"] = languages
        auto_match_all_images()
        set_progress("Data ready.", 100)

        return {
            "success": True,
            "total": len(SESSION["rows"]),
            "rows": SESSION["rows"],
            "languages": languages,
            "audio_url": "/api/stream_audio",
            "has_logo": bool(SESSION["logo_path"])
        }
    except Exception as e:
        return JSONResponse({"success": False, "error": str(e)}, status_code=500)


@app.post("/api/switch_format")
async def switch_format(payload: Dict[str, Any]):
    format_type = payload.get("format", "landscape")
    if SESSION.get("excel_path") and os.path.exists(SESSION["excel_path"]):
        rows, languages = parse_excel_sync_sheet(Path(SESSION["excel_path"]), format_type=format_type)
        if rows:
            SESSION["rows"] = rows
            SESSION["available_languages"] = languages
            auto_match_all_images()
            return {"success": True, "reloaded": True, "rows": SESSION["rows"], "languages": languages}
    return {"success": True, "reloaded": False, "rows": SESSION["rows"], "languages": SESSION["available_languages"]}


@app.post("/api/translate_timeline")
async def translate_timeline(payload: Dict[str, Any]):
    target_lang = payload.get("target_lang", "my")
    for r in SESSION["rows"]:
        translations = r.get("translations", {})
        if target_lang in translations and translations[target_lang]:
            r["caption"] = translations[target_lang]
        elif target_lang == "my":
            r["caption"] = r.get("original_text", r["caption"])
        r["target_lang"] = target_lang
    return {"success": True, "rows": SESSION["rows"]}


@app.post("/api/update_row")
async def update_row(payload: Dict[str, Any]):
    idx = int(payload.get("index", 0))
    if 0 <= idx < len(SESSION["rows"]):
        row = SESSION["rows"][idx]
        row["caption"] = payload.get("caption", row["caption"])
        row["start_time"] = float(payload.get("start_time", row["start_time"]))
        row["end_time"] = float(payload.get("end_time", row["end_time"]))
        row["duration"] = round(row["end_time"] - row["start_time"], 2)
        if "matched_img" in payload:
            row["matched_img"] = payload["matched_img"]
            row["is_manual_img"] = True
        return {"success": True, "row": row}
    return {"success": False}


@app.post("/api/delete_row")
async def delete_row(payload: Dict[str, Any]):
    idx = int(payload.get("index", -1))
    if 0 <= idx < len(SESSION["rows"]):
        SESSION["rows"].pop(idx)
        return {"success": True, "rows": SESSION["rows"]}
    return {"success": False, "error": "Invalid index."}


@app.post("/api/start_render")
async def start_render(
    background_tasks: BackgroundTasks,
    format: str = Form("landscape"),
    quality: str = Form("720p"),
    font_size: int = Form(44),
    line_spacing: int = Form(26),
    opacity: int = Form(45),
    enable_bgm: bool = Form(True),
    target_lang: str = Form("my"),
    rows_json: Optional[str] = Form(None)
):
    if not SESSION["rows"] or not SESSION["audio_path"]:
        return JSONResponse({"success": False, "error": "Missing assets."}, status_code=400)

    selected_lang = (target_lang or "my").strip().lower()

    # Use the complete rows sent by the browser so all preview edits are preserved.
    # In particular, do not restore captions from the original translations map.
    if rows_json:
        try:
            client_rows = json.loads(rows_json)
            if isinstance(client_rows, list) and client_rows:
                SESSION["rows"] = client_rows
        except (TypeError, json.JSONDecodeError):
            pass

    for row in SESSION["rows"]:
        row["target_lang"] = selected_lang

    SESSION["last_video"] = ""
    SESSION["is_rendering"] = True
    set_progress(f"Initializing [{selected_lang.upper()}] render...", 5)

    def _render_task():
        df = pd.DataFrame(SESSION["rows"])
        out_file = OUTPUT_DIR / f"dhamma_{format}_{quality}.mp4"

        cfg = {
            "format": format,
            "quality": quality,
            "font_size": font_size,
            "line_spacing": line_spacing,
            "opacity": opacity,
            "enable_bgm": enable_bgm,
            "target_lang": selected_lang
        }

        try:
            success, res = video_engine.render_all_clips(
                df=df,
                media_dir=SESSION["media_dir"],
                audio_path=SESSION["audio_path"],
                output_file=str(out_file),
                logo_path=SESSION["logo_path"],
                cfg=cfg,
                progress_cb=set_progress
            )
            if success and out_file.exists():
                SESSION["last_video"] = str(out_file)
                set_progress(f"[{selected_lang.upper()}] Ready! Click Download.", 100)
            else:
                set_progress(f"Render failed: {res}", 0)
        except Exception as exc:
            set_progress(f"Render failed: {exc}", 0)
        finally:
            SESSION["is_rendering"] = False

    background_tasks.add_task(_render_task)
    return {"success": True}


@app.get("/api/download_video")
async def download_video():
    if SESSION["last_video"] and os.path.exists(SESSION["last_video"]):
        filename = f"dhamma_video.mp4"
        return FileResponse(
            SESSION["last_video"],
            media_type="video/mp4",
            filename=filename,
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    return JSONResponse({"error": "Video not ready."}, status_code=404)


@app.get("/api/get_image/{filename}")
async def get_image(filename: str):
    safe_name = Path(filename).name
    for cand_dir in [Path(SESSION["media_dir"]), UPLOAD_DIR / "images", BASE_DIR / "images", BASE_DIR]:
        f_path = cand_dir / safe_name
        if f_path.is_file() and f_path.suffix.lower() in {".jpg", ".jpeg", ".png", ".webp"}:
            return FileResponse(f_path)
    return JSONResponse({"error": "Image not found"}, status_code=404)


@app.get("/api/stream_audio")
async def stream_audio():
    if os.path.exists(SESSION["audio_path"]):
        return FileResponse(SESSION["audio_path"], media_type="audio/mpeg")
    return JSONResponse({"error": "No audio loaded"}, status_code=404)


@app.get("/api/get_logo")
async def get_logo():
    if SESSION["logo_path"] and os.path.exists(SESSION["logo_path"]):
        return FileResponse(SESSION["logo_path"])
    return JSONResponse({"error": "No logo"}, status_code=404)


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("server:app", host="0.0.0.0", port=8000, reload=True)